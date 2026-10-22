#!/usr/bin/env python3
"""Build one design document as a .docx (and a .pdf when LibreOffice is installed) with its
figures embedded where the prose puts them, and the repository links pointing at GitHub.

Usage: python3 docs/tools/build_doc_docx.py docs/25-the-rulebook-and-the-board.md [out.docx]
       (default: build/<name>.docx and build/<name>.pdf)

Self-validating: the script opens the result, checks that every heading of the source is in
it, that every figure the source references is embedded, that no table was lost, and that no
em dash slipped in. Needs pandoc and python-docx; the .pdf needs soffice (LibreOffice).
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_master_docx import ROOT, load, reference_docx  # noqa: E402

AUTHOR = "Yobie Benjamin · yobie AT ieee DOT org · linkedin.com/in/yobie"


def landscape_reference() -> str:
    """The master reference document turned on its side: the tables in these documents are wide
    and the figures are drawn 10.4 inches across, so landscape pages keep both readable."""
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt
    path = reference_docx()
    d = docx.Document(path)
    for sec in d.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)   # US Letter on its side; pandoc's reference leaves the size unset
        sec.left_margin = sec.right_margin = Inches(0.6)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
    for name, size in (("Compact", 8.5), ("Body Text", 10), ("First Paragraph", 10)):
        try:
            d.styles[name].font.size = Pt(size)
        except KeyError:
            pass
    out = os.path.join(tempfile.gettempdir(), "sg-reference-landscape.docx")
    d.save(out)
    return out


def build(src_rel: str, out: str) -> str:
    text = open(os.path.join(ROOT, src_rel), encoding="utf-8").read()
    first = text.split("\n", 1)[0]
    title = first.lstrip("# ").strip() if first.startswith("# ") else os.path.basename(src_rel)
    md = load(src_rel, None)
    md = md.split("\n", 1)[1] if md.startswith("# ") else md          # the H1 becomes the title block
    head = [
        "---",
        f'title: "{title}"',
        f'author: "{AUTHOR}"',
        f'date: "{dt.date.today().isoformat()} · github.com/YobieBenjamin/safety-governor/blob/main/{src_rel}"',
        "---",
        "",
    ]
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as fh:
        fh.write("\n".join(head) + md)
        tmp = fh.name
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    # pandoc's markdown reader, not gfm: only it turns the separator row's dash counts into column
    # widths and makes a table span the text width (the gfm reader gives every column the same
    # width and leaves the table narrow, which is what made the first build unreadable)
    subprocess.run(["pandoc", tmp, "--from", "markdown+pipe_tables+raw_attribute+yaml_metadata_block+autolink_bare_uris+gfm_auto_identifiers-smart", "--to", "docx",
                    "--reference-doc", landscape_reference(), "--resource-path", os.path.join(ROOT, "docs"), "-o", out], check=True)
    os.unlink(tmp)
    return out


def fit_pictures(out: str) -> None:
    """pandoc scales a picture to the text width but not to the page height, so a tall figure
    runs off the page; this scales any such picture down to fit, keeping its proportions."""
    import docx
    d = docx.Document(out)
    sec = d.sections[0]
    max_w = sec.page_width - sec.left_margin - sec.right_margin
    max_h = sec.page_height - sec.top_margin - sec.bottom_margin - int(0.6 * 914400)   # room for the caption
    for shape in d.inline_shapes:
        w, h = shape.width, shape.height
        scale = min(max_w / w, max_h / h, 1.0)
        if scale < 1.0:
            shape.width, shape.height = int(w * scale), int(h * scale)
    d.save(out)


def validate(src_rel: str, out: str) -> None:
    import docx
    text = open(os.path.join(ROOT, src_rel), encoding="utf-8").read()
    d = docx.Document(out)
    body = "\n".join(p.text for p in d.paragraphs)
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    want = [re.sub(r"^#+\s*", "", ln).strip() for ln in text.split("\n") if re.match(r"^#{2,3}\s", ln)]
    missing = [w for w in want if not any(w[:40] in h for h in heads)]
    if missing:
        raise SystemExit(f"headings missing from the document: {missing}")
    figs = re.findall(r"\]\((?:docs/)?figures/(fig-[a-z-]+\.png)\)", text)
    n_img = sum(1 for rel in d.part.rels.values() if "image" in rel.reltype)
    if n_img < len(figs):
        raise SystemExit(f"expected {len(figs)} embedded figures, found {n_img}")
    n_src_tables = sum(1 for i, ln in enumerate(text.split("\n")) if ln.startswith("|") and i + 1 < len(text.split("\n")) and re.match(r"^\|?\s*:?-{3,}", text.split("\n")[i + 1]))
    if len(d.tables) < n_src_tables:
        raise SystemExit(f"expected {n_src_tables} tables, found {len(d.tables)}")
    if "—" in body:
        raise SystemExit("an em dash made it into the document")
    print(f"wrote {out}: {len(heads)} headings, {len(d.tables)} tables, {n_img} figures, {len(body.split()):,} words, {os.path.getsize(out) // 1024} KB")


def to_pdf(out: str) -> str | None:
    if not shutil.which("soffice"):
        print("no soffice: the .pdf was not made")
        return None
    outdir = os.path.dirname(os.path.abspath(out))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, out], check=True, capture_output=True)
    pdf = os.path.splitext(os.path.abspath(out))[0] + ".pdf"
    if not os.path.exists(pdf):
        raise SystemExit("soffice ran but produced no .pdf")
    print(f"wrote {pdf}: {os.path.getsize(pdf) // 1024} KB")
    return pdf


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = os.path.relpath(os.path.abspath(sys.argv[1]), ROOT)
    name = os.path.splitext(os.path.basename(src))[0]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "build", name + ".docx")
    build(src, out)
    fit_pictures(out)
    validate(src, out)
    to_pdf(out)
